"""Well master data and app settings.

The well table mirrors the 'Configuration Sheet' of Allocation.xlsm, one row per
well. Settings hold the values that lived in loose cells of the workbook
(conversion factors in C16:C18, tolerance in R5, the 200-iteration cap in the
Allocate macro, etc.).
"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
CONFIG_DIR = ROOT / "config"
WELLS_CSV = CONFIG_DIR / "wells.csv"            # well list loaded from your Configuration Sheet (git-ignored)
SETTINGS_JSON = CONFIG_DIR / "settings.json"

# column -> (label shown in the UI, workbook column on 'Configuration Sheet ')
WELL_COLUMNS: dict[str, tuple[str, str]] = {
    "well": ("Well", "C"),
    "uz_icv_tag": ("UZ ICV", "D"),
    "bhp_tag": ("UZ Tubing Pressure (psia)", "E"),
    "uz_annulus_tag": ("UZ Annulus Pressure (psia)", "F"),
    "oil_tag": ("MPFM Oil Rate (stb/d)", "G"),
    "gas_tag": ("MPFM Gas Rate (mmscf/d)", "H"),
    "water_tag": ("MPFM Water Rate (stb/d)", "I"),
    "uz_pi": ("UZ PI (stb/d/psi)", "J"),
    "mz_icv_tag": ("MZ ICV", "K"),
    "mz_tubing_tag": ("MZ Tubing Pressure (psia)", "L"),
    "mz_annulus_tag": ("MZ Annulus Pressure (psia)", "M"),
    "mz_pi": ("MZ PI (stb/d/psi)", "N"),
    "lz_icv_tag": ("LZ ICV", "O"),
    "lz_tubing_tag": ("LZ Tubing Pressure (psia)", "P"),
    "lz_annulus_tag": ("LZ Annulus Pressure (psia)", "Q"),
    "thp_tag": ("THP (psia)", "R"),
    "lz_pi": ("LZ PI (stb/d/psi)", "S"),
    "nodal_string": ("Nodal Analysis Pressure String", "T"),
    "liquid_factor": ("MPFM Liquid Rate Factor", "U"),
    "gor_factor": ("MPFM GOR Factor", "V"),
    "wct_factor": ("MPFM WCT Factor", "W"),
    "template": ("Zonal Sheet Template", "X"),
    "prosper_file": ("Prosper File Directory", "Y"),
    "split_uz": ("Zonal Split UZ", "Z"),
    "split_mz": ("Zonal Split MZ", "AA"),
    "split_lz": ("Zonal Split LZ", "AB"),
    "zone_uz": ("Zones UZ", "AC"),
    "zone_mz": ("Zones MZ", "AD"),
    "zone_lz": ("Zones LZ", "AE"),
}
NUMERIC_COLS = ["uz_pi", "mz_pi", "lz_pi", "liquid_factor", "gor_factor", "wct_factor",
                "split_uz", "split_mz", "split_lz"]

DEFAULT_SETTINGS: dict = {
    "mode": "live",                 # "live" (PI + PROSPER) or "demo"
    "pi_server": "",                # blank = default PI server
    "timezone": "Africa/Accra",
    "pi_interval": "1d",
    "liquid_conv": 151.0,           # m3/h -> stb/d   (workbook C16)
    "gas_conv": 0.000848,           # Sm3/h -> MMscf/d (workbook C17)
    "pressure_conv": 14.5,          # bar -> psi       (workbook C18)
    "tolerance": 0.05,              # fractional BHP mismatch (workbook R5)
    "max_iterations": 200,          # per phase (Allocate macro)
    "match_method": "sequential",   # "sequential" (workbook) or "binary_chop"
    "chop_min_frac": 0.1,           # binary chop rate bounds, x MPFM liquid
    "chop_max_frac": 3.0,
    "export_source": "mpfm",        # mpfm | matched | manual | auto
    "auto_threshold_pct": 5.0,      # auto: use MPFM when nodal deviation <= this
    "show_prosper": False,
}


def clean_text(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = str(v).strip()
    return "" if s in {"-", "nan", "None"} else s


def load_settings() -> dict:
    s = dict(DEFAULT_SETTINGS)
    if SETTINGS_JSON.exists():
        try:
            s.update(json.loads(SETTINGS_JSON.read_text()))
        except json.JSONDecodeError:
            pass
    return s


def save_settings(s: dict) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    SETTINGS_JSON.write_text(json.dumps(s, indent=2))


def normalise_wells(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in WELL_COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[list(WELL_COLUMNS)]
    for c in WELL_COLUMNS:
        if c in NUMERIC_COLS:
            df[c] = pd.to_numeric(df[c], errors="coerce")
        else:
            df[c] = df[c].map(clean_text)
    df["prosper_file"] = df["prosper_file"].str.strip('"')
    df = df[df["well"] != ""].reset_index(drop=True)
    return df


def _is_header_or_blank(first_cell) -> bool:
    v = clean_text(first_cell).lower()
    return v in {"", "well"}


def read_wells_csv(src) -> pd.DataFrame:
    """Read a well CSV (a path or an uploaded file) in either layout.

    * App layout: one header row with the WELL_COLUMNS keys (what save_wells writes).
    * Sheet layout: the Configuration Sheet saved as CSV from column C onwards,
      header rows included (Upper Zone / UZ ICV ... / units). Columns are taken by
      position C..AE, exactly like wells_from_workbook.
    """
    data = src.read() if hasattr(src, "read") else Path(src).read_bytes()
    text = data.decode("utf-8-sig", errors="replace") if isinstance(data, bytes) else data
    raw = pd.read_csv(io.StringIO(text), dtype=str, header=None, keep_default_na=False)
    if not raw.empty and "well" in {clean_text(v) for v in raw.iloc[0]}:
        return normalise_wells(pd.read_csv(io.StringIO(text), dtype=str))
    keys = list(WELL_COLUMNS)
    rows = []
    for _, r in raw.iterrows():
        if _is_header_or_blank(r.iloc[0]):
            continue
        vals = list(r.values)[:len(keys)]
        vals += [None] * (len(keys) - len(vals))
        rows.append(dict(zip(keys, vals)))
    return normalise_wells(pd.DataFrame(rows, columns=keys))


def load_wells() -> pd.DataFrame:
    """The loaded well list, or an empty table until a Configuration Sheet is attached."""
    if WELLS_CSV.exists():
        return read_wells_csv(WELLS_CSV)
    return normalise_wells(pd.DataFrame(columns=list(WELL_COLUMNS)))


def save_wells(df: pd.DataFrame) -> None:
    CONFIG_DIR.mkdir(parents=True, exist_ok=True)
    normalise_wells(df).to_csv(WELLS_CSV, index=False)


def wells_from_workbook(path_or_buffer) -> pd.DataFrame:
    """Read the 'Configuration Sheet ' of an Allocation.xlsm into the well table."""
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    wb = load_workbook(path_or_buffer, data_only=True, read_only=True, keep_vba=False)
    name = next((n for n in wb.sheetnames if n.strip().lower() == "configuration sheet"), None)
    if name is None:
        raise ValueError("No 'Configuration Sheet' found in the workbook.")
    ws = wb[name]
    rows = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        rec = {}
        for key, (_, col) in WELL_COLUMNS.items():
            idx = column_index_from_string(col) - 1
            rec[key] = r[idx] if idx < len(r) else None
        if clean_text(rec["well"]):
            rows.append(rec)
    return normalise_wells(pd.DataFrame(rows))


@dataclass(frozen=True)
class Zone:
    key: str          # UZ / MZ / LZ
    name: str         # zone label from the config sheet
    split: float      # allocation factor
    label: str        # column label used in the roll-up: "<well> <zone>"


def zones_for(well_row: pd.Series) -> list[Zone]:
    """Zones in play for a well, following templates Z1 (UZ), Z2 (UZ+LZ), Z3 (UZ+MZ+LZ)."""
    template = clean_text(well_row.get("template")).upper() or "Z1"
    keys = {"Z1": ["uz"], "Z2": ["uz", "lz"], "Z3": ["uz", "mz", "lz"]}.get(template, ["uz"])
    well = well_row["well"]
    out = []
    for k in keys:
        split = well_row.get(f"split_{k}")
        split = 1.0 if (template == "Z1" and (split is None or pd.isna(split))) else split
        if split is None or pd.isna(split):
            continue
        name = clean_text(well_row.get(f"zone_{k}")) or (well if template == "Z1" else k.upper())
        label = name if (name == well or name.startswith(well + " ")) else f"{well} {name}"
        out.append(Zone(k.upper(), name, float(split), label))
    return out
