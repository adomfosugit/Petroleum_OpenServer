"""Excel output laid out like the workbook's own sheets.

well_sheet       the per-well sheet ExportToTemplateSheet creates from Z1/Z2/Z3
month_workbook   every exported well + 'Gradient Match All wells' + 'Gradient Normalization' + 'PPRS'
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="DCE6EA")
ZONE = PatternFill("solid", fgColor="EEF3F5")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="B8C4CA")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

NUM_FMT = {"liquid": "#,##0", "oil": "#,##0", "water": "#,##0", "gas": "0.000",
           "wct": "0.00", "gor": "#,##0", "bhp_dev_pct": "0.00"}


def _cell(ws, row, col, value, fmt=None, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=None if (isinstance(value, float) and pd.isna(value)) else value)
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = BOLD
    if fill:
        c.fill = fill
    c.border = BOX
    return c


def well_sheet(ws, well: str, export: pd.DataFrame, meta: dict) -> None:
    """Template layout: B3 well name, row 4 headers, data from row 5; zone blocks from column K."""
    ws["B3"] = well
    ws["B3"].font = Font(bold=True, size=13)
    ws["B2"] = f"Template {meta.get('template', '')}, rates: {meta.get('source_label', '')}"
    ws["B2"].font = Font(italic=True, color="5A6B74")
    base = [("Date", None), ("Liquid Rate", "liquid"), ("WCT", "wct"), ("GOR", "gor"),
            ("BHP Percentage deviation", "bhp_dev_pct"), ("Oil Rate", "oil"), ("Gas Rate", "gas"),
            ("Water Rate", "water")]
    for j, (h, _) in enumerate(base):
        _cell(ws, 4, 2 + j, h, bold=True, fill=HEAD).alignment = CENTER
    zones = meta.get("zones", [])
    col = 11  # column K, as on the templates
    if zones:
        ws.cell(row=2, column=10, value="Factor Based")
    for z in zones:
        ws.cell(row=1, column=col, value=z["name"]).font = BOLD
        ws.cell(row=2, column=col, value=z["split"])
        ws.cell(row=3, column=col, value=z["key"]).font = BOLD
        for k, h in enumerate(["Oil", "Water", "Gas", "GOR", "WCT"]):
            _cell(ws, 4, col + k, h, bold=True, fill=ZONE).alignment = CENTER
        col += 5
    for i, (d, r) in enumerate(export.iterrows()):
        row = 5 + i
        _cell(ws, row, 2, d.to_pydatetime(), "dd-mmm-yyyy")
        for j, (_, key) in enumerate(base[1:]):
            _cell(ws, row, 3 + j, r.get(key), NUM_FMT.get(key))
        col = 11
        for z in zones:
            for k, (ph, fmt) in enumerate([("oil", "#,##0"), ("water", "#,##0"), ("gas", "0.000"),
                                           ("gor", "#,##0"), ("wct", "0.000")]):
                _cell(ws, row, col + k, r.get(f"{z['label']}|{ph}"), fmt)
            col += 5
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 13
    for c in range(3, col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12
    ws.freeze_panes = "C5"


def _matrix_block(ws, top: int, title: str, m: pd.DataFrame, fmt: str):
    """Block layout of the all-wells sheets: title, well row, header row, data rows (dates in column F)."""
    ws.cell(row=top, column=6, value=title).font = Font(bold=True, size=12)
    for j, c in enumerate(m.columns):
        ws.cell(row=top + 1, column=7 + j, value=c.split(" ")[0] if c not in ("Total", "PPRS") else None).font = BOLD
        _cell(ws, top + 2, 7 + j, c, bold=True, fill=HEAD).alignment = CENTER
    _cell(ws, top + 2, 6, "Date", bold=True, fill=HEAD)
    for i, (d, r) in enumerate(m.iterrows()):
        _cell(ws, top + 3 + i, 6, d.to_pydatetime(), "dd-mmm-yyyy")
        for j, v in enumerate(r.values):
            _cell(ws, top + 3 + i, 7 + j, float(v) if pd.notna(v) else None, fmt)
    ws.column_dimensions["F"].width = 13
    for j in range(len(m.columns)):
        ws.column_dimensions[get_column_letter(7 + j)].width = 12


# workbook rows: oil header row 3 (data row 4), gas header row 62 (data 63), water header row 117 (data 118)
BLOCK_TOPS = {"oil": 1, "gas": 60, "water": 115}
FMT = {"oil": "#,##0", "gas": "0.000", "water": "#,##0"}


def well_workbook(well: str, export: pd.DataFrame, meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = well[:31]
    well_sheet(ws, well, export, meta)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_workbook(exports: dict, matrices: dict, normalised: dict, pprs: pd.DataFrame) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    gm = wb.create_sheet("Gradient Match All wells")
    for ph, m in matrices.items():
        _matrix_block(gm, BLOCK_TOPS[ph], ph.title(), m, FMT[ph])
    gm.freeze_panes = "G4"
    gn = wb.create_sheet("Gradient Normalization")
    for ph, m in normalised.items():
        _matrix_block(gn, BLOCK_TOPS[ph], ph.title(), m, FMT[ph])
    gn.freeze_panes = "G4"
    if not pprs.empty:
        ps = wb.create_sheet("PPRS")
        from .store import PPRS_COLUMNS
        _cell(ps, 2, 5, "Date", bold=True, fill=HEAD)
        for j, (k, label) in enumerate(PPRS_COLUMNS.items()):
            _cell(ps, 2, 6 + j, label, bold=True, fill=HEAD).alignment = CENTER
            ps.column_dimensions[get_column_letter(6 + j)].width = 16
        for i, (d, r) in enumerate(pprs.iterrows()):
            _cell(ps, 3 + i, 5, d.to_pydatetime(), "dd-mmm-yyyy")
            for j, k in enumerate(PPRS_COLUMNS):
                _cell(ps, 3 + i, 6 + j, float(r[k]) if pd.notna(r[k]) else None, "#,##0.00")
        ps.column_dimensions["E"].width = 13
    for well, (df, meta) in exports.items():
        well_sheet(wb.create_sheet(well[:31]), well, df, meta)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
