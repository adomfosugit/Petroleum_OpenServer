import openpyxl
import pytest

from petroleum_openserver.exceptions import ProgramTerminated
from petroleum_openserver.gap_validation import get_flowline_inlet_joint_index, set_well_routing


@pytest.fixture
def workbook():
    wb = openpyxl.Workbook()
    config = wb.active
    config.title = "Configuration"
    config["A1"] = "Well Name"
    for i, (well, gap) in enumerate([("A01", "A01"), ("B02", "B02")], start=3):
        config[f"A{i}"] = well
        config[f"B{i}"] = gap

    ws = wb.create_sheet("Validation")
    ws["A1"] = "Well Name"
    ws["B1"] = "Riser"
    ws["A5"] = "A01"
    ws["B5"] = "R11"
    ws["A6"] = "B02"
    ws["B6"] = "R10"
    return wb


def test_set_well_routing_sets_valve_states(fake_client, workbook):
    set_well_routing(fake_client, workbook["Configuration"], workbook["Validation"])

    assert ("GAP.MOD[{PROD}].VALVE[{A01_a}].StateValue", 0) in fake_client.set_calls
    assert ("GAP.MOD[{PROD}].VALVE[{A01_b}].StateValue", 1) in fake_client.set_calls
    assert ("GAP.MOD[{PROD}].VALVE[{B02_a}].StateValue", 1) in fake_client.set_calls
    assert ("GAP.MOD[{PROD}].VALVE[{B02_b}].StateValue", 0) in fake_client.set_calls


def test_set_well_routing_rejects_unknown_riser(fake_client, workbook):
    workbook["Validation"]["B6"] = "R99"
    with pytest.raises(ProgramTerminated):
        set_well_routing(fake_client, workbook["Configuration"], workbook["Validation"])


def test_get_flowline_inlet_joint_index_finds_matching_joint(fake_client):
    fake_client.get_values["GAP.MOD[0].JOINT.COUNT"] = 3
    fake_client.get_values["GAP.MOD[0].JOINT[2].ISCONNECTEDTO(MOD[0].PIPE[{A01_fl}])"] = 1
    fake_client.get_values["GAP.MOD[0].JOINT[2].ISCONNECTEDTO(MOD[0].INLGEN[{A01_chk}])"] = 1

    assert get_flowline_inlet_joint_index(fake_client, "A01") == 2


def test_get_flowline_inlet_joint_index_defaults_to_zero_when_not_found(fake_client):
    fake_client.get_values["GAP.MOD[0].JOINT.COUNT"] = 2
    assert get_flowline_inlet_joint_index(fake_client, "A01") == 0
