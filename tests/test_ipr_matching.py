import openpyxl
import pytest

from petroleum_openserver.ipr import ipr_rate
from petroleum_openserver.ipr_matching import estimate_res_press


@pytest.fixture
def workbook():
    wb = openpyxl.Workbook()
    config = wb.active
    config.title = "Configuration"
    config["A1"] = "Bubble Pt"
    config["C1"] = 2200.0  # Offset(0, 2) from "Bubble Pt"

    ws = wb.create_sheet("Validation")
    for col, text in {
        "A": "Well Name",
        "B": "Number of Zones",
        "C": "UZ Liquid Rate",
        "D": "UZ PI",
        "E": "UZ Watercut",
        "F": "UZ Annulus Pressure",
        "G": "UZ Tubing Pressure",
        "H": "UZ ICV",
        "I": "UZ ICV dP",
        "J": "UZ Res. Pres.",
        "K": "PROSPER BHP",
    }.items():
        ws[f"{col}1"] = text

    ws["A5"] = "TESTWELL"
    ws["B5"] = 1
    return wb


def test_estimate_res_press_uses_annulus_pressure_when_available(workbook):
    ws = workbook["Validation"]
    pi, pb, wc, pwf_gauge = 1.8, 2200.0, 15.0, 1200.0
    true_pr_gauge = 4300.0
    q = ipr_rate(true_pr_gauge, pwf_gauge, pi, pb, wc)

    ws["C5"] = q
    ws["D5"] = pi
    ws["E5"] = wc
    ws["F5"] = pwf_gauge + 14.7  # annulus pressure is stored as absolute (psia)

    estimate_res_press(workbook["Configuration"], ws, 0, "UZ")

    assert ws["J5"].value == pytest.approx(true_pr_gauge + 14.7, abs=0.1)


def test_estimate_res_press_leaves_existing_value_when_no_valid_bhp(workbook):
    ws = workbook["Validation"]
    ws["C5"] = 1500
    ws["D5"] = 1.5
    ws["E5"] = 10
    ws["H5"] = 5  # ICV open
    ws["J5"] = 4000  # pre-existing user-entered reservoir pressure

    estimate_res_press(workbook["Configuration"], ws, 0, "UZ")

    assert ws["J5"].value == 4000


def test_estimate_res_press_flags_error_when_no_bhp_and_no_prior_value(workbook):
    ws = workbook["Validation"]
    ws["C5"] = 1500
    ws["D5"] = 1.5
    ws["E5"] = 10
    ws["H5"] = 5  # ICV open, so this isn't "closed" - should flag ERROR

    estimate_res_press(workbook["Configuration"], ws, 0, "UZ")

    assert ws["J5"].value == "ERROR"


def test_estimate_res_press_skips_when_not_producing(workbook):
    ws = workbook["Validation"]
    ws["C5"] = 0  # zone not producing

    estimate_res_press(workbook["Configuration"], ws, 0, "UZ")

    assert ws["J5"].value is None
