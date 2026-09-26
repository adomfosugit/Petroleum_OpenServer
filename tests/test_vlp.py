import openpyxl
import pytest

from petroleum_openserver.vlp import vlp_calc


@pytest.fixture
def ws():
    wb = openpyxl.Workbook()
    sheet = wb.active

    headers = {
        "A": "Well Name",
        "B": "WHP",
        "C": "Watercut",
        "D": "GOR",
        "E": "Liquid Rate",
        "F": "UZ Tubing Pressure",
        "G": "PROSPER BHP",
        "H": "IPR dP Offset",
        "I": "Number of Zones",
        "J": "UZ IPR dP Offset",
        "K": "MZ IPR dP Offset",
        "L": "LZ IPR dP Offset",
    }
    for col, text in headers.items():
        sheet[f"{col}1"] = text

    # Field unit == sheet unit throughout, so ipt_unit_conv is a no-op and the
    # test can assert on plain numbers without a real conversion table.
    for col, unit in {"B": "psi", "C": "pct", "D": "scf/stb", "E": "stb/d", "F": "psi", "G": "psi"}.items():
        sheet[f"{col}3"] = unit
        sheet[f"{col}4"] = unit

    sheet["A5"] = "TESTWELL"
    sheet["B5"] = 500
    sheet["C5"] = 20
    sheet["D5"] = 800
    sheet["E5"] = 3000
    sheet["F5"] = 1000  # measured UZ tubing (gauge) pressure
    sheet["I5"] = 1  # single-zone well

    return sheet


def test_vlp_calc_sets_prosper_inputs(fake_client, ws):
    fake_client.get_values["PROSPER.OUT.VLP.RES[0][0].PRES"] = 1200

    vlp_calc(fake_client, ws, 0)

    assert ("PROSPER.ANL.VLP.Pres", 500.0) in fake_client.set_calls
    assert ("PROSPER.ANL.VLP.WC", 20.0) in fake_client.set_calls
    assert ("PROSPER.ANL.VLP.GORVAL", 800.0) in fake_client.set_calls
    assert ("PROSPER.ANL.VLP.Rates[0]", 3000.0) in fake_client.set_calls
    assert "PROSPER.ANL.VLP.CALC" in fake_client.cmds


def test_vlp_calc_writes_bhp_and_ipr_offset(fake_client, ws):
    fake_client.get_values["PROSPER.OUT.VLP.RES[0][0].PRES"] = 1200

    vlp_calc(fake_client, ws, 0)

    assert ws["G5"].value == 1200
    assert ws["H5"].value == 200  # calculated (1200) minus measured (1000)
    assert ws["J5"].value == 200  # single zone -> UZ IPR dP Offset mirrors it


def test_vlp_calc_skips_when_inputs_invalid(fake_client, ws):
    ws["B5"] = "-"  # WHP not numeric -> nothing should be sent to PROSPER

    vlp_calc(fake_client, ws, 0)

    assert fake_client.set_calls == []
    assert ws["G5"].value == "-"
    assert ws["H5"].value == "-"
